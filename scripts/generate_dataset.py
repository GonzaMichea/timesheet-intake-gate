"""Genera el maestro de empleados y 6 lotes sintéticos con defectos sembrados.

Uso: python scripts/generate_dataset.py [--seed 42]

El ground truth sale del plan de siembra (lo que este script rompe a propósito), nunca de
ejecutar el motor. La disposición esperada se calcula aquí con una implementación propia y
simple de las reglas del contrato.
"""
import argparse
import csv
import io
import json
import random
import re
import shutil
import sys
import zipfile
from datetime import date, datetime, timedelta
from pathlib import Path

import yaml
from openpyxl import Workbook

ROOT = Path(__file__).resolve().parents[1]
SPEC = yaml.safe_load((ROOT / "spec" / "timesheet_spec.yaml").read_text(encoding="utf-8"))
COLUMNS = list(SPEC["columns"])
SITES = SPEC["sites"]
PERIOD = "2026-09"
FIRST = ["Camila", "Matías", "Valentina", "Joaquín", "Fernanda", "Tomás", "Javiera", "Benjamín",
         "Antonia", "Vicente", "Catalina", "Cristóbal", "Daniela", "Sebastián", "Francisca"]
LAST = ["Soto", "Muñoz", "Rojas", "Díaz", "Pérez", "Vega", "Castro", "Reyes", "Fuentes", "Silva",
        "Araya", "Tapia", "Mora", "Lagos", "Núñez"]
NOTES = ["revisar", "turno extra", "reemplazo", "ok", "pendiente"]
FIXED_ZIP_DATE = (2026, 10, 1, 0, 0, 0)


# ---------- maestro ----------
def build_master(rng: random.Random) -> dict[str, list[dict]]:
    """60 empleados: 15 por sede; los 2 últimos de cada sede están inactivos."""
    master = {}
    for si, site in enumerate(SITES, start=1):
        staff = []
        for n in range(1, 16):
            staff.append({"id": f"E{si}{n:04d}", "name": f"{rng.choice(FIRST)} {rng.choice(LAST)}",
                          "active": n <= 13})
        master[site] = staff
    return master


def write_master(master) -> None:
    path = ROOT / "spec" / "employee_master.csv"
    with path.open("w", encoding="utf-8", newline="") as handle:
        writer = csv.writer(handle, lineterminator="\n")
        writer.writerow(["employee_id", "site", "active"])
        for site in SITES:
            for emp in master[site]:
                writer.writerow([emp["id"], site, "true" if emp["active"] else "false"])


# ---------- datos base ----------
def weekdays(period: str) -> list[date]:
    year, month = map(int, period.split("-"))
    day, days = date(year, month, 1), []
    while day.month == month:
        if day.weekday() < 5:
            days.append(day)
        day += timedelta(days=1)
    return days


def base_rows(site: str, staff: list[dict], rng: random.Random) -> list[dict]:
    active = [e for e in staff if e["active"]]
    rows = []
    for day in weekdays(PERIOD):
        for index, emp in enumerate(active):
            absent = rng.random() < 0.03
            rows.append({
                "employee_id": emp["id"], "employee_name": emp["name"], "site": site,
                "cost_center": f"CC{100 + index % 4}", "work_date": day.isoformat(),
                "shift_code": ["DAY", "DAY", "NIGHT", "SWING"][index % 4],
                "regular_hours": "0" if absent else "8",
                "overtime_hours": "0" if absent else rng.choice(["0", "0", "1", "2"]),
                "absence_code": rng.choice(["VAC", "SICK", "PERM"]) if absent else "",
                "approved_by": f"Supervisor {site}",
                "approval_date": f"2026-10-0{rng.randint(1, 3)}",
            })
    return rows


# ---------- disposición esperada (implementación propia, a partir del YAML) ----------
def severity(code: str, column: str) -> str:
    if code == "D07":
        return SPEC["columns"][column].get("d07_severity", "ERROR")
    return SPEC["severities"][code]


def expected_disposition(units: list[tuple], total_rows: int) -> str:
    sev = [(severity(code, col), code, row) for code, col, row in units]
    if any(s == "BLOCKER" for s, _, _ in sev):
        return "BLOCK"
    errors = [(c, r) for s, c, r in sev if s == "ERROR"]
    if errors:
        full = any(c in SPEC["file_level_full_impact"] for c, _ in errors)
        pct = 1.0 if full else len({r for _, r in errors}) / total_rows
        if round(pct - SPEC["thresholds"]["quarantine_error_row_pct"], 9) > 0:
            return "QUARANTINE"
    return "LOAD_WITH_WARNINGS" if units else "LOAD"


# ---------- constructor de archivos con defectos sembrados ----------
class FileBuilder:
    def __init__(self, site, master, rng, fmt="csv", sep=",", name=None):
        self.site, self.rng, self.fmt, self.sep = site, rng, fmt, sep
        self.name = name or f"TS_{site}_{PERIOD.replace('-', '')}.{fmt}"
        self.staff = master[site]
        self.rows = base_rows(site, self.staff, rng)
        self.units: list[tuple] = []  # (code, columna, fila)  fila 0 = defecto de archivo/columna
        self.mappings: list[dict] = []
        self.renames: dict[str, str] = {}
        self.dropped: set[str] = set()
        self.extras: list[str] = []
        self.used: set[int] = set()
        self.dup_count = 0

    def limit(self, n):
        self.rows = self.rows[:n]
        return self

    def pick(self, n, distinct_dates=False):
        order = list(range(len(self.rows)))
        self.rng.shuffle(order)
        chosen, dates = [], set()
        for i in order:
            date_value = self.rows[i]["work_date"]
            if i in self.used or (distinct_dates and date_value in dates):
                continue
            chosen.append(i)
            dates.add(date_value)
            if len(chosen) == n:
                break
        self.used.update(chosen)
        return chosen

    def cell(self, code, column, n, value, rule_column=None, distinct_dates=False, extra=None):
        for k, i in enumerate(self.pick(n, distinct_dates)):
            self.rows[i][column] = value(k, i) if callable(value) else value
            for other, other_value in (extra or {}).items():
                self.rows[i][other] = other_value
            self.units.append((code, column if rule_column is None else rule_column, i + 2))
        return self

    # D04 + D05: columna renombrada (la plantilla antigua) con su destino esperado
    def rename(self, column, header):
        self.renames[column] = header
        self.units += [("D04", column, 0), ("D05", header, 0)]
        self.mappings.append({"file": self.name, "column": header, "target": column})
        return self

    def drop(self, column):  # D04 sin candidata
        self.dropped.add(column)
        self.units.append(("D04", column, 0))
        return self

    def extra(self, header):  # D06
        self.extras.append(header)
        self.units.append(("D06", header, 0))
        return self

    def file_defect(self, code):  # D02 / D03
        self.units.append((code, "", 0))
        return self

    def duplicates(self, n):  # D08: copias de filas limpias, agregadas al final
        free = [i for i in range(len(self.rows)) if i not in self.used]
        sources = self.rng.sample(free, n)
        self.used.update(sources)
        first_new = len(self.rows)
        for k, i in enumerate(sources):
            self.rows.append(dict(self.rows[i]))
            self.units.append(("D08", "", first_new + k + 2))
        return self

    def headers(self):
        base = [self.renames.get(c, c) for c in COLUMNS if c not in self.dropped]
        return base + self.extras

    def table(self):
        out = []
        for row in self.rows:
            record = {self.renames.get(c, c): row[c] for c in COLUMNS if c not in self.dropped}
            for header in self.extras:
                record[header] = self.rng.choice(NOTES)
            out.append(record)
        return out

    def content(self) -> bytes:
        headers, table = self.headers(), self.table()
        if self.fmt == "csv":
            buffer = io.StringIO()
            writer = csv.DictWriter(buffer, fieldnames=headers, delimiter=self.sep, lineterminator="\n")
            writer.writeheader()
            writer.writerows(table)
            return buffer.getvalue().encode("utf-8")
        return xlsx_bytes(headers, table)

    def disposition(self):
        return expected_disposition(self.units, len(self.rows))


def xlsx_bytes(headers, table) -> bytes:
    workbook = Workbook()
    workbook.properties.creator = "generate_dataset"
    workbook.properties.created = workbook.properties.modified = datetime(2026, 10, 1)
    sheet = workbook.active
    sheet.append(headers)
    for record in table:
        cells = []
        for header in headers:
            value = record[header]
            try:  # las fechas ISO se guardan como fechas reales de Excel
                value = datetime.strptime(value, "%Y-%m-%d") if len(value) == 10 and value[4] == "-" else value
            except ValueError:
                pass
            cells.append(value)
        sheet.append(cells)
    raw = io.BytesIO()
    workbook.save(raw)
    # Se reescribe el zip con fecha fija para que el archivo sea idéntico byte a byte
    fixed = io.BytesIO()
    with zipfile.ZipFile(raw) as src, zipfile.ZipFile(fixed, "w", zipfile.ZIP_DEFLATED) as dst:
        for info in sorted(src.infolist(), key=lambda i: i.filename):
            entry = zipfile.ZipInfo(info.filename, date_time=FIXED_ZIP_DATE)
            entry.compress_type = zipfile.ZIP_DEFLATED
            data = src.read(info)
            if info.filename == "docProps/core.xml":  # openpyxl pone aquí la hora actual al guardar
                data = re.sub(rb"(<dcterms:modified[^>]*>)[^<]*", rb"\g<1>2026-10-01T00:00:00Z", data)
            dst.writestr(entry, data)
    return fixed.getvalue()


# ---------- lotes ----------
class Batch:
    def __init__(self, name):
        self.name, self.files, self.entries, self.defects, self.mappings = name, {}, [], [], []

    def add(self, builder: FileBuilder):
        self.files[builder.name] = builder.content()
        self.entries.append({"file": builder.name, "site": builder.site,
                             "expected_disposition": builder.disposition()})
        for code, column, row in builder.units:
            self.defects.append((builder.name, code, column, row))
        self.mappings += builder.mappings

    def junk(self, name, content=b"sede;horas\nSCL;100\n"):  # D02
        self.files[name] = content
        self.entries.append({"file": name, "site": None, "expected_disposition": "QUARANTINE"})
        self.defects.append((name, "D02", "", 0))

    def missing(self, site):  # D01
        key = f"MISSING:{site}"
        self.entries.append({"file": key, "site": site, "expected_disposition": "BLOCK"})
        self.defects.append((key, "D01", "", 0))

    def truth(self):
        grouped: dict[tuple, list[int]] = {}
        for file, code, column, row in self.defects:
            grouped.setdefault((file, code, column), [])
            if row:
                grouped[(file, code, column)].append(row)
        return {
            "batch": self.name, "period": PERIOD, "files": self.entries,
            "defects": [{"file": f, "code": c, "column": col, "rows": sorted(rows)}
                        for (f, c, col), rows in sorted(grouped.items())],
            "mappings": self.mappings,
        }


def make_batches(master, seed):
    def rng(batch, site):
        return random.Random(f"{seed}-{batch}-{site}")

    def builder(batch, site, **kwargs):
        return FileBuilder(site, master, rng(batch, site), **kwargs)

    batches = []

    b = Batch("batch_00_clean")
    for site, fmt in zip(SITES, ["csv", "xlsx", "csv", "xlsx"]):
        b.add(builder(b.name, site, fmt=fmt))
    batches.append(b)

    b = Batch("batch_01_old_template")  # SCL exporta con plantilla antigua; VAP no envía
    b.add(builder(b.name, "SCL").rename("overtime_hours", "Horas Extra").rename("cost_center", "Centro Costo")
          .extra("Notas").duplicates(3).cell("D09", "work_date", 2, lambda k, i: f"2026-08-{k + 1:02d}"))
    b.missing("VAP")
    b.add(builder(b.name, "CCP"))
    b.add(builder(b.name, "ANF", fmt="xlsx"))
    batches.append(b)

    b = Batch("batch_02_wrong_files")  # archivo de agosto y un resumen suelto
    wrong = builder(b.name, "SCL", name="TS_SCL_202608.csv").file_defect("D03")
    b.add(wrong)
    b.junk("resumen_horas.csv")
    b.add(builder(b.name, "VAP", sep=";")
          .cell("D07", "regular_hours", 2, "ocho").cell("D10", "regular_hours", 2, "13", extra={"overtime_hours": "0"})
          .cell("D11", "shift_code", 2, "NOCHE"))
    b.add(builder(b.name, "CCP")
          .cell("D10", "regular_hours", 3, "12", rule_column="", extra={"overtime_hours": "3"})
          .cell("D11", "absence_code", 2, "LIC"))
    b.add(builder(b.name, "ANF", fmt="xlsx"))
    batches.append(b)

    b = Batch("batch_03_approvals")  # aprobaciones incompletas y personal dado de baja
    inactive = [e["id"] for e in master["SCL"] if not e["active"]]
    b.add(builder(b.name, "SCL", sep=";")
          .cell("D12", "employee_id", 3, lambda k, i: f"E9{k + 1:04d}", distinct_dates=True)
          .cell("D12", "employee_id", 4, lambda k, i: inactive[k % 2], distinct_dates=True)
          .cell("D12", "approved_by", 3, ""))
    b.add(builder(b.name, "VAP", fmt="xlsx")
          .cell("D07", "work_date", 3, "31/31/2026").cell("D12", "approval_date", 2, ""))
    b.add(builder(b.name, "CCP").drop("approved_by"))
    b.add(builder(b.name, "ANF", fmt="xlsx").extra("Comentarios"))
    batches.append(b)

    b = Batch("batch_04_partial")  # faltan sedes y un nombre en minúsculas
    b.missing("SCL")
    b.missing("VAP")
    b.junk("ts_ccp_202609.csv", builder(b.name, "CCP").content())
    b.missing("CCP")
    b.add(builder(b.name, "ANF", fmt="xlsx").rename("shift_code", "Turno")
          .cell("D10", "overtime_hours", 2, "5", extra={"regular_hours": "8"}).duplicates(2))
    batches.append(b)

    b = Batch("batch_05_mixed")  # mezcla para cubrir bordes
    b.add(builder(b.name, "SCL", name="TS_SCL_202610.csv").file_defect("D03"))
    b.add(builder(b.name, "VAP").rename("employee_id", "Employee ID").extra("Obs")
          .cell("D09", "work_date", 2, lambda k, i: f"2026-10-{k + 1:02d}"))
    b.add(builder(b.name, "CCP").limit(250).cell("D11", "shift_code", 5, "TURNO"))  # 5/250 = 2 % exacto
    inactive_anf = [e["id"] for e in master["ANF"] if not e["active"]]
    b.add(builder(b.name, "ANF", fmt="xlsx")
          .cell("D12", "employee_id", 2, lambda k, i: inactive_anf[k], distinct_dates=True)
          .cell("D07", "overtime_hours", 1, "2h").cell("D09", "work_date", 1, "2026-08-15"))
    batches.append(b)
    return batches


def write_batches(batches, out: Path):
    if out.exists():
        shutil.rmtree(out)
    for batch in batches:
        folder = out / batch.name
        folder.mkdir(parents=True)
        for name, content in batch.files.items():
            (folder / name).write_bytes(content)
        text = json.dumps(batch.truth(), ensure_ascii=False, indent=2) + "\n"
        (folder / "ground_truth.json").write_bytes(text.encode("utf-8"))  # siempre LF, también en Windows


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__.splitlines()[0])
    parser.add_argument("--seed", type=int, default=42)
    parser.add_argument("--out", default=str(ROOT / "datasets"))
    args = parser.parse_args()
    master = build_master(random.Random(f"{args.seed}-master"))
    write_master(master)
    batches = make_batches(master, args.seed)
    write_batches(batches, Path(args.out))
    for batch in batches:
        print(f"{batch.name}: {len(batch.files)} archivo(s), {len(batch.defects)} defecto(s) sembrado(s)")
    return 0


if __name__ == "__main__":
    sys.exit(main())
