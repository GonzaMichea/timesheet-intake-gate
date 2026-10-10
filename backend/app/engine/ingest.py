"""Lectura de archivos: CSV, XLSX y ZIP. Todo se lee como texto (R2.3–R2.5)."""
import csv
import io
import zipfile
from dataclasses import dataclass
from datetime import date, datetime
from pathlib import PurePosixPath

from openpyxl import load_workbook

from backend.app.engine.models import InputFile


class IngestError(Exception):
    """El archivo no se pudo leer; el mensaje está en español."""


@dataclass
class Table:
    headers: list[str]
    records: list[dict[str, str]]
    row_nums: list[int]  # numeración Excel: encabezado = 1, primer dato = 2


def expand_inputs(files: list[tuple[str, bytes]]) -> list[InputFile]:
    """Descomprime los .zip y los trata como archivos individuales (R2.3)."""
    result: list[InputFile] = []
    for name, content in files:
        if not name.lower().endswith(".zip"):
            result.append(InputFile(name=name, content=content))
            continue
        try:
            with zipfile.ZipFile(io.BytesIO(content)) as archive:
                for info in archive.infolist():
                    parts = PurePosixPath(info.filename.replace("\\", "/")).parts
                    base = parts[-1] if parts else ""
                    if info.is_dir() or "__MACOSX" in parts or base.startswith("."):
                        continue
                    result.append(InputFile(name=base, content=archive.read(info)))
        except zipfile.BadZipFile as error:
            raise IngestError(f"El archivo {name} está dañado y no se pudo abrir.") from error
    return result


def check_total_size(files: list[tuple[str, bytes]], max_mb: float) -> None:
    if sum(len(c) for _, c in files) > max_mb * 1024 * 1024:
        raise IngestError(f"La carga supera el máximo de {max_mb:g} MB.")


def _cell_text(value) -> str:
    if value is None:
        return ""
    if isinstance(value, datetime):
        return value.strftime("%Y-%m-%d")
    if isinstance(value, date):
        return value.isoformat()
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value)


def _build_table(rows: list[list[str]]) -> Table:
    if not rows or not any(cell.strip() for cell in rows[0]):
        raise IngestError("El archivo no tiene fila de encabezados.")
    headers = [cell.strip() for cell in rows[0]]
    records: list[dict[str, str]] = []
    row_nums: list[int] = []
    for index, row in enumerate(rows[1:], start=2):
        if not any(cell.strip() for cell in row):
            continue  # fila vacía: se conserva la numeración de las demás
        padded = list(row) + [""] * (len(headers) - len(row))
        records.append(dict(zip(headers, padded[: len(headers)])))
        row_nums.append(index)
    return Table(headers=headers, records=records, row_nums=row_nums)


def _read_csv(content: bytes) -> Table:
    try:
        text = content.decode("utf-8-sig")
    except UnicodeDecodeError:
        text = content.decode("latin-1")
    first_line = text.splitlines()[0] if text.strip() else ""
    delimiter = ";" if first_line.count(";") > first_line.count(",") else ","
    return _build_table(list(csv.reader(io.StringIO(text), delimiter=delimiter)))


def _read_xlsx(content: bytes) -> Table:
    try:
        workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
        sheet = workbook.worksheets[0]
        rows = [[_cell_text(v) for v in row] for row in sheet.iter_rows(values_only=True)]
    except Exception as error:
        raise IngestError("El archivo Excel no se pudo leer.") from error
    return _build_table(rows)


def read_table(name: str, content: bytes) -> Table:
    if name.lower().endswith(".xlsx"):
        return _read_xlsx(content)
    return _read_csv(content)
